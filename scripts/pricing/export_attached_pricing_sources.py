#!/usr/bin/env python3
"""Normalize heterogeneous courier pricing sources (Excel + PDF) into import-ready CSV files.

This script is intentionally adapter-based:
- Excel adapter: reads city-level rate rows.
- PDF adapter: extracts city-to-zone mapping from zone list tables.

Outputs are written to an export directory so they can be reviewed and then imported into
Laravel pricing settings (zone master, city map, and lane/rate seeds).
"""

from __future__ import annotations

import argparse
import csv
import json
import re
from collections import Counter, defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Dict, Iterable, List, Optional, Sequence, Set, Tuple

from openpyxl import load_workbook

try:
    import pdfplumber
except ImportError:  # pragma: no cover
    pdfplumber = None


HEADER_ALIASES = {
    "couriercode": "courier_code",
    "courier_code": "courier_code",
    "countryname": "country",
    "province_name": "province",
    "district_name": "district",
    "city_name": "city",
    "firstkg": "first_kg",
    "first_kg": "first_kg",
    "extra": "extra_kg",
    "extrakg": "extra_kg",
    "extra_kg": "extra_kg",
}

ZONE_RE = re.compile(r"^zone\s*([a-z0-9]+)$", flags=re.IGNORECASE)
SPACE_RE = re.compile(r"\s+")
CITY_KEY_RE = re.compile(r"[^a-z0-9]+")


@dataclass
class ExcelRateRow:
    company_code: str
    country: str
    province: str
    district: str
    city: str
    zone_code: str
    first_kg: str
    extra_kg: str


@dataclass
class CityZoneRow:
    company_code: str
    service_category: str
    country: str
    city: str
    zone_code: str
    source: str
    province: str = ""
    district: str = ""


def normalize_header(header: object) -> str:
    if header is None:
        return ""
    normalized = re.sub(r"[^a-z0-9]+", "", str(header).strip().lower())
    return HEADER_ALIASES.get(normalized, normalized)


def normalize_space(value: object) -> str:
    if value is None:
        return ""
    return SPACE_RE.sub(" ", str(value).strip())


def normalize_zone_code(value: object) -> str:
    raw = normalize_space(value)
    if not raw:
        return ""
    match = ZONE_RE.match(raw)
    if match:
        return match.group(1).upper()
    return raw.upper()


def parse_decimal(value: object) -> str:
    if value is None:
        return ""
    text = str(value).strip().replace(",", "")
    if not text:
        return ""
    try:
        number = Decimal(text)
    except InvalidOperation:
        return text
    if number == number.to_integral_value():
        return str(number.quantize(Decimal("1")))
    return format(number.normalize(), "f")


def city_key(city_name: str) -> str:
    lowered = normalize_space(city_name).lower()
    return CITY_KEY_RE.sub("", lowered)


def build_header_index(header_row: Sequence[object]) -> Dict[str, int]:
    index: Dict[str, int] = {}
    for i, header in enumerate(header_row):
        normalized = normalize_header(header)
        if normalized and normalized not in index:
            index[normalized] = i
    return index


def read_excel_rates(excel_path: Path, default_country: str) -> List[ExcelRateRow]:
    workbook = load_workbook(excel_path, data_only=True, read_only=True)
    sheet = workbook["courier city rates"] if "courier city rates" in workbook.sheetnames else workbook.active

    iterator = sheet.iter_rows(values_only=True)
    try:
        header_row = next(iterator)
    except StopIteration:
        return []

    header_index = build_header_index(header_row)
    required = ["courier_code", "city", "zone", "first_kg", "extra_kg"]
    missing = [column for column in required if column not in header_index]
    if missing:
        raise ValueError(f"Missing required Excel columns: {', '.join(missing)}")

    rows: List[ExcelRateRow] = []
    for row in iterator:
        company_code = normalize_space(row[header_index["courier_code"]])
        city = normalize_space(row[header_index["city"]])
        zone_code = normalize_zone_code(row[header_index["zone"]])
        if not city or not zone_code:
            continue

        country = normalize_space(row[header_index.get("country", -1)]) if "country" in header_index else ""
        province = normalize_space(row[header_index.get("province", -1)]) if "province" in header_index else ""
        district = normalize_space(row[header_index.get("district", -1)]) if "district" in header_index else ""

        rows.append(
            ExcelRateRow(
                company_code=company_code,
                country=(country or default_country),
                province=province,
                district=district,
                city=city,
                zone_code=zone_code,
                first_kg=parse_decimal(row[header_index["first_kg"]]),
                extra_kg=parse_decimal(row[header_index["extra_kg"]]),
            )
        )

    return rows


def split_cell_lines(cell: object) -> List[str]:
    text = normalize_space(cell)
    if not text:
        return []

    raw = str(cell).replace("\r", "\n")
    values: List[str] = []
    for line in raw.split("\n"):
        cleaned = normalize_space(line)
        if not cleaned:
            continue
        if cleaned.lower().startswith("zone "):
            continue
        values.append(cleaned)

    if values:
        return values
    return [text]


def find_zone_column(header_row: Sequence[object], label: str) -> Optional[int]:
    target = label.lower()
    for i, cell in enumerate(header_row):
        if normalize_space(cell).lower() == target:
            return i
    return None


def parse_pdf_zones(pdf_path: Path) -> Dict[str, Set[str]]:
    if pdfplumber is None:
        raise RuntimeError(
            "pdfplumber is required for PDF parsing. Install it in the active Python environment first."
        )

    zones: Dict[str, Set[str]] = defaultdict(set)

    with pdfplumber.open(str(pdf_path)) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                if not table:
                    continue

                header_index: Optional[int] = None
                for row_index, row in enumerate(table):
                    row_text = " | ".join(normalize_space(value) for value in row if value)
                    if "Zone A" in row_text and "Zone B" in row_text and "Zone C" in row_text:
                        header_index = row_index
                        break

                if header_index is None:
                    continue

                header_row = table[header_index]
                zone_a_col = find_zone_column(header_row, "Zone A")
                zone_b_col = find_zone_column(header_row, "Zone B")
                zone_c_col = find_zone_column(header_row, "Zone C")
                zone_d_col = find_zone_column(header_row, "Zone D")

                if zone_c_col is None:
                    continue

                if zone_d_col is not None and zone_d_col > zone_c_col:
                    zone_c_columns = list(range(zone_c_col, zone_d_col))
                else:
                    zone_c_columns = [zone_c_col]

                for row in table[header_index + 1 :]:
                    if zone_a_col is not None and not zones.get("A"):
                        if zone_a_col < len(row):
                            zone_a_values = split_cell_lines(row[zone_a_col])
                            if zone_a_values:
                                zones["A"].add(zone_a_values[0])

                    if zone_b_col is not None and zone_b_col < len(row):
                        for name in split_cell_lines(row[zone_b_col]):
                            zones["B"].add(name)

                    for c_col in zone_c_columns:
                        if c_col >= len(row):
                            continue
                        for name in split_cell_lines(row[c_col]):
                            zones["C"].add(name)

                    if zone_d_col is not None and zone_d_col < len(row):
                        for name in split_cell_lines(row[zone_d_col]):
                            zones["D"].add(name)

    return zones


def write_csv(path: Path, fieldnames: Sequence[str], rows: Iterable[Dict[str, object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(handle, fieldnames=fieldnames)
        writer.writeheader()
        for row in rows:
            writer.writerow(row)


def build_excel_city_rate_rows(
    rates: Sequence[ExcelRateRow],
    company_code: str,
    service_category: str,
    currency: str,
    source_file: str,
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for item in rates:
        rows.append(
            {
                "company_code": company_code,
                "service_category": service_category,
                "country": item.country,
                "province": item.province,
                "district": item.district,
                "city": item.city,
                "zone_code": item.zone_code,
                "first_kg": item.first_kg,
                "extra_kg": item.extra_kg,
                "currency": currency,
                "source_file": source_file,
            }
        )
    return rows


def build_excel_zone_rate_rows(
    rates: Sequence[ExcelRateRow],
    company_code: str,
    service_category: str,
    currency: str,
    source_file: str,
) -> List[Dict[str, str]]:
    totals = Counter(item.zone_code for item in rates)
    grouped: Dict[Tuple[str, str, str], int] = Counter((item.zone_code, item.first_kg, item.extra_kg) for item in rates)

    rows: List[Dict[str, str]] = []
    for (zone_code, first_kg, extra_kg), city_count_for_rate in sorted(grouped.items()):
        rows.append(
            {
                "company_code": company_code,
                "service_category": service_category,
                "country": rates[0].country if rates else "",
                "zone_code": zone_code,
                "first_kg": first_kg,
                "extra_kg": extra_kg,
                "currency": currency,
                "city_count_for_rate": str(city_count_for_rate),
                "total_cities_in_zone": str(totals.get(zone_code, 0)),
                "source_file": source_file,
            }
        )
    return rows


def build_pdf_city_zone_rows(
    zones: Dict[str, Set[str]],
    company_code: str,
    service_category: str,
    country: str,
    source_file: str,
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for zone_code in sorted(zones):
        for city in sorted(zones[zone_code]):
            rows.append(
                {
                    "company_code": company_code,
                    "service_category": service_category,
                    "country": country,
                    "city": city,
                    "zone_code": zone_code,
                    "source_file": source_file,
                }
            )
    return rows


def build_pdf_zone_master_rows(
    zones: Dict[str, Set[str]],
    company_code: str,
    service_category: str,
    country: str,
    source_file: str,
) -> List[Dict[str, str]]:
    rows: List[Dict[str, str]] = []
    for zone_code in sorted(zones):
        rows.append(
            {
                "company_code": company_code,
                "service_category": service_category,
                "country": country,
                "zone_code": zone_code,
                "zone_label": f"Zone {zone_code}",
                "city_count": str(len(zones[zone_code])),
                "is_catch_all": "1" if "restofthecities" in {city_key(v) for v in zones[zone_code]} else "0",
                "source_file": source_file,
            }
        )
    return rows


def build_merged_city_zone_rows(
    excel_rows: Sequence[ExcelRateRow],
    pdf_zones: Dict[str, Set[str]],
    company_code: str,
    service_category: str,
    country: str,
    excel_source: str,
    pdf_source: str,
) -> Tuple[List[Dict[str, str]], List[Dict[str, str]]]:
    merged: List[Dict[str, str]] = []
    city_zone_sources: Dict[str, Set[Tuple[str, str]]] = defaultdict(set)

    for item in excel_rows:
        row = CityZoneRow(
            company_code=company_code,
            service_category=service_category,
            country=item.country or country,
            city=item.city,
            zone_code=item.zone_code,
            source=excel_source,
            province=item.province,
            district=item.district,
        )
        merged.append(row.__dict__)
        city_zone_sources[city_key(item.city)].add((item.zone_code, excel_source))

    for zone_code, cities in pdf_zones.items():
        for city in sorted(cities):
            row = CityZoneRow(
                company_code=company_code,
                service_category=service_category,
                country=country,
                city=city,
                zone_code=zone_code,
                source=pdf_source,
            )
            merged.append(row.__dict__)
            city_zone_sources[city_key(city)].add((zone_code, pdf_source))

    conflicts: List[Dict[str, str]] = []
    for key, zone_source_set in sorted(city_zone_sources.items()):
        zones = sorted({zone for zone, _source in zone_source_set})
        if len(zones) <= 1:
            continue

        sources = sorted({source for _zone, source in zone_source_set})
        conflicts.append(
            {
                "city_key": key,
                "zones": "|".join(zones),
                "sources": "|".join(sources),
                "assignments": "|".join(f"{zone}:{source}" for zone, source in sorted(zone_source_set)),
            }
        )

    return merged, conflicts


def resolve_company_code(rows: Sequence[ExcelRateRow], explicit: Optional[str]) -> str:
    if explicit:
        return explicit
    discovered = sorted({row.company_code for row in rows if row.company_code})
    if len(discovered) == 1:
        return discovered[0]
    if discovered:
        return discovered[0]
    return "unknown_company"


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Normalize attached courier pricing files into canonical CSV exports.")
    parser.add_argument("--excel", required=True, type=Path, help="Path to Excel city-rate file")
    parser.add_argument("--pdf", required=True, type=Path, help="Path to PDF zone-list file")
    parser.add_argument("--out-dir", required=True, type=Path, help="Directory for normalized exports")
    parser.add_argument("--service-category", default="domestic", help="Service category label")
    parser.add_argument("--country", default="LK", help="Country code")
    parser.add_argument("--currency", default="LKR", help="Currency code")
    parser.add_argument("--company-code", default=None, help="Override company code")
    return parser.parse_args()


def main() -> None:
    args = parse_args()

    excel_rows = read_excel_rates(args.excel, default_country=args.country)
    if not excel_rows:
        raise ValueError("No rate rows parsed from Excel input.")

    pdf_zones = parse_pdf_zones(args.pdf)
    if not pdf_zones:
        raise ValueError("No zones parsed from PDF input.")

    company_code = resolve_company_code(excel_rows, args.company_code)

    excel_source = args.excel.name
    pdf_source = args.pdf.name

    city_rate_rows = build_excel_city_rate_rows(
        excel_rows,
        company_code=company_code,
        service_category=args.service_category,
        currency=args.currency,
        source_file=excel_source,
    )
    zone_rate_rows = build_excel_zone_rate_rows(
        excel_rows,
        company_code=company_code,
        service_category=args.service_category,
        currency=args.currency,
        source_file=excel_source,
    )
    pdf_city_zone_rows = build_pdf_city_zone_rows(
        pdf_zones,
        company_code=company_code,
        service_category=args.service_category,
        country=args.country,
        source_file=pdf_source,
    )
    pdf_zone_master_rows = build_pdf_zone_master_rows(
        pdf_zones,
        company_code=company_code,
        service_category=args.service_category,
        country=args.country,
        source_file=pdf_source,
    )
    merged_city_zone_rows, conflict_rows = build_merged_city_zone_rows(
        excel_rows,
        pdf_zones,
        company_code=company_code,
        service_category=args.service_category,
        country=args.country,
        excel_source=excel_source,
        pdf_source=pdf_source,
    )

    out_dir = args.out_dir
    out_dir.mkdir(parents=True, exist_ok=True)

    write_csv(
        out_dir / "city_rate_card_from_excel.csv",
        [
            "company_code",
            "service_category",
            "country",
            "province",
            "district",
            "city",
            "zone_code",
            "first_kg",
            "extra_kg",
            "currency",
            "source_file",
        ],
        city_rate_rows,
    )

    write_csv(
        out_dir / "zone_rate_card_from_excel.csv",
        [
            "company_code",
            "service_category",
            "country",
            "zone_code",
            "first_kg",
            "extra_kg",
            "currency",
            "city_count_for_rate",
            "total_cities_in_zone",
            "source_file",
        ],
        zone_rate_rows,
    )

    write_csv(
        out_dir / "city_zone_map_from_pdf.csv",
        [
            "company_code",
            "service_category",
            "country",
            "city",
            "zone_code",
            "source_file",
        ],
        pdf_city_zone_rows,
    )

    write_csv(
        out_dir / "zone_master_from_pdf.csv",
        [
            "company_code",
            "service_category",
            "country",
            "zone_code",
            "zone_label",
            "city_count",
            "is_catch_all",
            "source_file",
        ],
        pdf_zone_master_rows,
    )

    write_csv(
        out_dir / "city_zone_map_merged.csv",
        [
            "company_code",
            "service_category",
            "country",
            "city",
            "zone_code",
            "source",
            "province",
            "district",
        ],
        merged_city_zone_rows,
    )

    write_csv(
        out_dir / "zone_conflicts.csv",
        ["city_key", "zones", "sources", "assignments"],
        conflict_rows,
    )

    summary = {
        "company_code": company_code,
        "excel_row_count": len(excel_rows),
        "excel_zone_count": len({row.zone_code for row in excel_rows}),
        "pdf_zone_count": len(pdf_zones),
        "pdf_city_count": sum(len(cities) for cities in pdf_zones.values()),
        "merged_city_zone_rows": len(merged_city_zone_rows),
        "conflict_count": len(conflict_rows),
        "output_files": [
            "city_rate_card_from_excel.csv",
            "zone_rate_card_from_excel.csv",
            "city_zone_map_from_pdf.csv",
            "zone_master_from_pdf.csv",
            "city_zone_map_merged.csv",
            "zone_conflicts.csv",
        ],
    }

    with (out_dir / "export_summary.json").open("w", encoding="utf-8") as handle:
        json.dump(summary, handle, indent=2)

    print(json.dumps(summary, indent=2))


if __name__ == "__main__":
    main()
